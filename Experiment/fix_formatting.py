import sys
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
import re

def fix_doc():
    doc_path = r'c:\Minor Project\Documentation\Research paper_updated.docx'
    doc = Document(doc_path)
    
    # 1. Remove Table 1 (Table II. EXPERIMENTAL ENVIRONMENT)
    if len(doc.tables) > 1:
        tbl = doc.tables[1]._element
        tbl.getparent().remove(tbl)
    
    # 2. Re-style all tables to "Normal Table"
    for t in doc.tables:
        t.style = 'Normal Table'
    
    # Track which paragraphs to move or delete
    p_to_delete = []
    p_to_move = []
    p_target = None
    
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        
        # 3. Find paragraphs to delete
        if "TABLE II. EXPERIMENTAL ENVIRONMENT" in text:
            p_to_delete.append(p)
        elif text == "VII. Methodology":
            p_to_delete.append(p)
            
        # 4. Text replacements
        if "(Table II)" in text:
            p.text = p.text.replace("(Table II)", "")
            p.text = re.sub(r' +', ' ', p.text) # fix double spaces
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            
        if "TABLE III. QUANTIZATION RESULTS" in text:
            p.text = "TABLE II. QUANTIZATION RESULTS — GEMMA 2 2B INSTRUCT"
            p.style = 'Normal'
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
        if "Table III summarizes" in text:
            p.text = p.text.replace("Table III", "Table II")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            
        if "TABLE IV. STATISTICAL PC BENCHMARKS" in text:
            p.text = "TABLE III. STATISTICAL PC BENCHMARKS (QWEN 2.5 1.5B, N=15)"
            p.style = 'Normal'
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
        if "Table IV summarizes" in text:
            p.text = p.text.replace("Table IV", "Table III")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            
        if "Table V: Android Emulator" in text:
            p.text = "TABLE IV. ANDROID EMULATOR (PIXEL 7 AVD) PERFORMANCE METRICS"
            p.style = 'Normal'
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
        # 5. Fix Subheadings
        if "C. Android Emulator Inference Performance" in text:
            p.text = p.text.replace("C. Android", "B. Android")
            p.style = 'Normal'
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            
        if "A. Statistically Controlled PC Benchmarks" in text:
            p.style = 'Normal'
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            
        if "A. Quantization Strategy (Layer-Wise" in text:
            p.style = 'Normal'
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p_to_move.append(p)
            
        if "To achieve optimal performance on mobile devices with limited RAM" in text:
            p.style = 'Normal'
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p_to_move.append(p)
            
        if "This methodology—developed and tested in our mixed-quantization" in text:
            p.style = 'Normal'
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p_to_move.append(p)
            
        # Target for moving
        if "VIII. EXPERIMENTAL SETUP AND PRELIMINARY RESULTS" in text:
            p_target = p
            
        # Maintain justification for normal text paragraphs (e.g. Discussion)
        if text.startswith("Compared with F16, the statistically controlled"):
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            
        if text.startswith("To simulate on-device execution"):
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            
        if text.startswith("Due to the emulator's memory constraints"):
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Move paragraphs
    if p_target and p_to_move:
        parent = p_target._element.getparent()
        target_idx = parent.index(p_target._element)
        for pm in p_to_move:
            # We must move the element, not copy it
            pm_element = pm._element
            parent.insert(target_idx, pm_element)
            target_idx += 1 # advance index so they stay in order
            
    # Delete paragraphs
    for pd in p_to_delete:
        pe = pd._element
        pe.getparent().remove(pe)
        
    doc.save(r'c:\Minor Project\Documentation\Research paper_final.docx')
    print("Done formatting!")

if __name__ == '__main__':
    fix_doc()
