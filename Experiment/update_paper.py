import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc_path = r'c:\Minor Project\Documentation\Research paper.docx'
doc = docx.Document(doc_path)

insert_idx = -1
for i, p in enumerate(doc.paragraphs):
    if 'Gemma 2 2B Instruct was converted' in p.text:
        insert_idx = i
        break

if insert_idx != -1:
    target_p = doc.paragraphs[insert_idx + 1]
    
    new_heading = target_p.insert_paragraph_before("A. Statistically Controlled PC Benchmarks (Qwen 2.5 1.5B)")
    new_p1 = target_p.insert_paragraph_before(
        "To move beyond single-run directional evidence, a rigorous benchmarking pipeline was implemented. "
        "The Qwen 2.5 1.5B Instruct model (quantized to Q4_K_M) was evaluated over 15 repeated, statistically controlled iterations "
        "on the development PC (CPU-only execution). This experiment measures load time, Time to First Token (TTFT), "
        "sustained generation speed (Tokens/sec), and peak RAM utilization."
    )
    new_p2 = target_p.insert_paragraph_before(
        "Table IV summarizes the findings, reporting the mean (μ), standard deviation (σ), and 95% confidence intervals (CI) "
        "across all 15 runs. The tight confidence intervals confirm that the Q4_K_M quantization provides stable and highly predictable "
        "latency and memory footprints, validating its selection for the constrained Android environment."
    )
    table_caption = target_p.insert_paragraph_before("TABLE IV. STATISTICAL PC BENCHMARKS (QWEN 2.5 1.5B Q4_K_M, N=15)")
    table_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Create table at the end of the document
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Metric'
    hdr_cells[1].text = 'Mean (μ)'
    hdr_cells[2].text = 'Std Dev (σ)'
    hdr_cells[3].text = '95% CI'
    
    data = [
        ("Model Load Time", "0.854 s", "0.088 s", "± 0.044 s"),
        ("Time to First Token (TTFT)", "0.244 s", "0.013 s", "± 0.007 s"),
        ("Tokens per Second", "36.72 tok/s", "1.29 tok/s", "± 0.65 tok/s"),
        ("Peak RAM Utilization", "1667.37 MB", "1.84 MB", "± 0.93 MB")
    ]
    for item in data:
        row_cells = table.add_row().cells
        row_cells[0].text = item[0]
        row_cells[1].text = item[1]
        row_cells[2].text = item[2]
        row_cells[3].text = item[3]
        
    # Move the table to right before target_p
    target_p._p.addprevious(table._tbl)
    
    doc.save(doc_path)
    print("Successfully updated the research paper with PC benchmarks.")
else:
    print("Could not find the target paragraph to insert the results.")
