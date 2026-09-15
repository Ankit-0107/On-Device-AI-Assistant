import json
import os
from docx import Document
from pathlib import Path

BASE_DIR = Path(r"c:\Minor Project")
RESULTS_PATH = BASE_DIR / "Experiment" / "tier1_pc_results_all.json"
DOC_PATH = BASE_DIR / "Documentation" / "Research paper.docx"

def update_paper():
    # Load JSON results
    with open(RESULTS_PATH, 'r', encoding='utf-8') as f:
        data = json.load(f)
        
    summary = data["summary"]
    
    # Load document
    doc = Document(str(DOC_PATH))
    
    # 1. Update Table IV (Index 3)
    t = doc.tables[3]
    # Clear existing rows except header
    while len(t.rows) > 1:
        t._element.remove(t.rows[-1]._element)
        
    # Update header
    hdr = t.rows[0].cells
    hdr[0].text = "Variant"
    hdr[1].text = "Tok/s (mean ± σ)"
    hdr[2].text = "95% CI"
    hdr[3].text = "TTFT (mean ± σ)"
    
    # If the table only had 4 columns, we'll add one more for Peak RAM, or just keep it 4 and merge RAM into a string
    # Actually, we can just add a cell. Let's just create a completely new table and replace the old one.
    
    # Let's delete the old table and insert a new one in its place.
    tbl_element = t._element
    parent = tbl_element.getparent()
    idx = parent.index(tbl_element)
    parent.remove(tbl_element)
    
    # Create new table
    new_tbl = doc.add_table(rows=1, cols=5)
    new_tbl.style = 'Table Grid'
    
    hdr_cells = new_tbl.rows[0].cells
    hdr_cells[0].text = "Variant"
    hdr_cells[1].text = "Tok/s (mean ± σ)"
    hdr_cells[2].text = "95% CI (Tok/s)"
    hdr_cells[3].text = "TTFT (mean ± σ)"
    hdr_cells[4].text = "Peak RAM"
    
    # Add rows from data
    order = ["Qwen2.5-1.5B-Instruct-f16.gguf", "Qwen2.5-1.5B-Instruct-Q8_0.gguf", "qwen2.5-1.5b-instruct-q4_k_m.gguf"]
    labels = ["F16", "Q8_0", "Q4_K_M"]
    
    for key, label in zip(order, labels):
        stats = summary[key]["metrics"]
        row = new_tbl.add_row().cells
        
        tok_mean = stats["tok_sec"]["mean"]
        tok_std = stats["tok_sec"]["std_dev"]
        tok_ci = stats["tok_sec"]["ci_95"]
        
        ttft_mean = stats["ttft_s"]["mean"]
        ttft_std = stats["ttft_s"]["std_dev"]
        
        ram_mean = stats["peak_ram_mb"]["mean"]
        
        row[0].text = label
        row[1].text = f"{tok_mean:.1f} ± {tok_std:.2f}"
        row[2].text = f"± {tok_ci:.2f}"
        row[3].text = f"{ttft_mean:.2f}s ± {ttft_std:.2f}s"
        row[4].text = f"{ram_mean:.0f} MB"

    # Move new table to correct position (where old table was)
    parent.insert(idx, new_tbl._element)

    # 2. Update Paragraphs
    # Note: Paragraph numbers might shift slightly if we do insertions, but here we just edit text in place.
    # We found them at 45, 46, and 49 earlier. Let's find by content to be safe.
    
    for p in doc.paragraphs:
        if "To move beyond single-run directional evidence, a rigorous benchmarking pipeline was implemented." in p.text:
            p.text = "To move beyond single-run directional evidence, a rigorous benchmarking pipeline was implemented. The Qwen 2.5 1.5B Instruct model was evaluated across its F16, Q8_0, and Q4_K_M formats over 15 repeated, statistically controlled iterations on the development PC (using CPU-only execution, matching the target mobile environment). This experiment measured Time to First Token (TTFT), sustained generation speed (Tokens/sec), and peak RAM utilization to directly compare the architectural tradeoffs of quantization."
            
        elif "Table IV summarizes the findings, reporting the mean" in p.text:
            p.text = "Table IV summarizes the findings, reporting the mean (μ), standard deviation (σ), and 95% confidence intervals (CI) for each quantization variant. The tight confidence intervals confirm that the performance gains from Q8_0 and Q4_K_M are statistically significant and provide highly predictable latency and memory footprints, validating the selection of Q4_K_M for the constrained Android environment."
            
        elif "Compared with F16, Q8_0 cut file size" in p.text:
            p.text = "Compared with F16, the statistically controlled multi-run CPU benchmarks show that Q8_0 drastically improved throughput (from ~12 to ~24 tok/s) and almost halved the RAM utilization (from ~3050 MB to ~1674 MB). Q4_K_M was smaller still, recording the highest throughput of the three (~35 tok/s), consistent with the literature's finding that decode is memory-bandwidth-bound on consumer hardware. Importantly, the tight 95% confidence intervals confirm that these performance gains are structurally robust, not artifacts of run-to-run variance. Q4_K_M remains the leading candidate for smartphone deployment because it fits remarkably well into a typical 4–8 GB mobile RAM budget while retaining high execution speed."
            
        elif "TABLE IV. STATISTICAL PC BENCHMARKS" in p.text:
            p.text = "TABLE IV. STATISTICAL PC BENCHMARKS (QWEN 2.5 1.5B, N=15)"

    doc.save(str(DOC_PATH.with_name("Research paper_updated.docx")))
    print("Paper updated successfully! (Saved as Research paper_updated.docx)")

if __name__ == "__main__":
    update_paper()
